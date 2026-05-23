"""
backend/app/services/report_service.py
─────────────────────────────────────────
PDF (fpdf2) and Excel (openpyxl) report generation for JEE Dropout Prediction System.
Both functions return raw bytes for FastAPI StreamingResponse.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from typing   import Any, Optional

log = logging.getLogger(__name__)


# ═══════════════════════════════════════════════════════════════════════════════
# PDF — Student Risk Report
# ═══════════════════════════════════════════════════════════════════════════════

def generate_student_pdf(
    student:            Any,
    latest_assessment:  Optional[Any],
    shap_explanation:   Optional[Any],
    mock_tests:         list[Any],
    institute_name:     str = "JEE Coaching Institute",
) -> bytes:
    """
    Generate a professionally styled PDF risk report for a single student.
    Returns raw PDF bytes. Raises RuntimeError if fpdf2 is unavailable.
    """
    try:
        from fpdf import FPDF, XPos, YPos
    except ImportError:
        raise RuntimeError("fpdf2 not installed. Run: pip install fpdf2")

    # ── Color constants ───────────────────────────────────────────────────────
    NAVY    = (15,  23,  42)
    INDIGO  = (99,  102, 241)
    WHITE   = (255, 255, 255)
    SLATE50 = (248, 250, 252)
    SLATE   = (100, 116, 139)
    TEXT    = (30,  41,  59)

    RISK_COLORS: dict[str, tuple] = {
        "Low":      (16,  185, 129),
        "Medium":   (245, 158, 11),
        "High":     (239, 68,  68),
        "Critical": (124, 58,  237),
    }

    risk_level = getattr(latest_assessment, "risk_level", "Unknown")
    risk_score = getattr(latest_assessment, "risk_score", 0) or 0
    risk_color = RISK_COLORS.get(risk_level, SLATE)

    # ── PDF setup ─────────────────────────────────────────────────────────────
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_margins(18, 10, 18)

    def _rgb(color): return color[0], color[1], color[2]

    def _safe(text: str) -> str:
        """Strip characters outside Latin-1 so fpdf2 Helvetica never raises."""
        return text.encode("latin-1", errors="replace").decode("latin-1")

    # ── 1. Header ─────────────────────────────────────────────────────────────
    pdf.set_fill_color(*NAVY)
    pdf.rect(0, 0, 210, 28, style="F")
    pdf.set_xy(18, 7)
    pdf.set_font("Helvetica", "B", 14)
    pdf.set_text_color(*WHITE)
    pdf.cell(140, 8, "JEE Dropout Risk Report", new_x=XPos.RIGHT, new_y=YPos.TOP)
    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(148, 163, 184)
    pdf.set_xy(18, 17)
    pdf.cell(0, 5, f"{institute_name}  |  Generated: {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}")
    pdf.ln(22)

    # ── 2. Student Info Box ───────────────────────────────────────────────────
    pdf.set_fill_color(*SLATE50)
    pdf.set_draw_color(226, 232, 240)
    y = pdf.get_y()
    pdf.rect(18, y, 174, 32, style="FD")
    pdf.set_xy(24, y + 5)
    pdf.set_font("Helvetica", "B", 11)
    pdf.set_text_color(*TEXT)
    pdf.cell(0, 6, getattr(student, "full_name", "Unknown Student"))
    pdf.set_xy(24, y + 13)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*SLATE)
    pdf.cell(60, 5, f"Code: {getattr(student, 'student_code', '-')}")
    pdf.cell(60, 5, f"Batch ID: {getattr(student, 'batch_id', '-')}")
    enroll = getattr(student, "enrollment_date", None)
    enroll_str = enroll.strftime("%d %b %Y") if enroll else "-"
    pdf.cell(0, 5, f"Enrolled: {enroll_str}")
    pdf.set_xy(24, y + 21)
    pdf.cell(60, 5, f"Target Rank: {getattr(student, 'target_rank', '-')}")
    pdf.cell(0, 5, f"Target College: {getattr(student, 'target_college', '-')}")
    pdf.ln(38)

    # ── 3. Risk Assessment Box ────────────────────────────────────────────────
    pdf.set_fill_color(*risk_color)
    y = pdf.get_y()
    pdf.rect(18, y, 174, 30, style="FD")
    pdf.set_xy(24, y + 5)
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(*WHITE)
    pdf.cell(30, 14, f"{risk_score:.0f}")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_xy(54, y + 6)
    pdf.cell(0, 5, "/ 100  RISK SCORE")
    pdf.set_xy(54, y + 14)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 5, f"{risk_level.upper()} RISK")

    # ASCII risk bar
    bar_filled = int(risk_score / 5)
    bar_empty  = 20 - bar_filled
    bar_str    = "#" * bar_filled + "." * bar_empty
    pdf.set_xy(110, y + 10)
    pdf.set_font("Courier", "", 8)
    pdf.cell(80, 5, bar_str)
    pdf.ln(36)

    # ── 4. SHAP Top Factors Table ─────────────────────────────────────────────
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(*TEXT)
    pdf.cell(0, 7, "Key Risk Factors (SHAP Analysis)", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_draw_color(226, 232, 240)

    top_factors: list[dict] = []
    if shap_explanation and hasattr(shap_explanation, "top_factors"):
        raw = shap_explanation.top_factors
        if isinstance(raw, list):
            top_factors = raw
        elif isinstance(raw, dict):
            top_factors = raw.get("factors", [])

    headers = ["Factor", "Actual Value", "SHAP Impact", "Direction"]
    widths  = [72, 36, 36, 30]

    pdf.set_fill_color(*NAVY)
    pdf.set_text_color(*WHITE)
    pdf.set_font("Helvetica", "B", 8)
    for h, w in zip(headers, widths):
        pdf.cell(w, 7, h, border=1, fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    for i, factor in enumerate(top_factors[:8]):
        pdf.set_fill_color(248, 250, 252) if i % 2 == 0 else pdf.set_fill_color(*WHITE)
        direction = factor.get("direction", "")
        arrow     = "[+] Increases" if direction == "increases_risk" else "[-] Decreases"
        pdf.set_text_color(*TEXT)
        pdf.cell(72, 6, str(factor.get("human_label", factor.get("feature", "-")))[:38], border=1, fill=True)
        pdf.cell(36, 6, str(factor.get("actual_value", "-")), border=1, fill=True)
        pdf.cell(36, 6, f"{factor.get('shap_value', 0):.4f}", border=1, fill=True)
        pdf.cell(30, 6, arrow, border=1, fill=True)
        pdf.ln()

    if not top_factors:
        pdf.set_text_color(*SLATE)
        pdf.cell(0, 6, "  No SHAP explanation available for this assessment.", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.ln(6)

    # ── 5. Mock Score History ────────────────────────────────────────────────
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(*TEXT)
    pdf.cell(0, 7, "Recent Mock Test History (Last 5)", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    last5 = sorted(mock_tests, key=lambda t: getattr(t, "test_date", datetime.min), reverse=True)[:5]
    t_headers = ["Test Name", "Date", "Total", "Physics", "Chemistry", "Maths"]
    t_widths  = [52, 28, 20, 22, 28, 24]

    pdf.set_fill_color(*NAVY)
    pdf.set_text_color(*WHITE)
    pdf.set_font("Helvetica", "B", 8)
    for h, w in zip(t_headers, t_widths):
        pdf.cell(w, 7, h, border=1, fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    for i, t in enumerate(last5):
        pdf.set_fill_color(248, 250, 252) if i % 2 == 0 else pdf.set_fill_color(*WHITE)
        pdf.set_text_color(*TEXT)
        dt = getattr(t, "test_date", None)
        dt_str = dt.strftime("%d %b %y") if dt else "—"
        pdf.cell(52, 6, str(getattr(t, "test_name", "—"))[:28], border=1, fill=True)
        pdf.cell(28, 6, dt_str,                                                     border=1, fill=True)
        pdf.cell(20, 6, f"{getattr(t, 'total_score',     0):.0f}",                  border=1, fill=True)
        pdf.cell(22, 6, f"{getattr(t, 'physics_score',   0):.0f}",                  border=1, fill=True)
        pdf.cell(28, 6, f"{getattr(t, 'chemistry_score', 0):.0f}",                  border=1, fill=True)
        pdf.cell(24, 6, f"{getattr(t, 'maths_score',     0):.0f}",                  border=1, fill=True)
        pdf.ln()

    if not last5:
        pdf.set_text_color(*SLATE)
        pdf.cell(0, 6, "  No mock test records found.", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.ln(6)

    # ── 6. Recommendations ────────────────────────────────────────────────────
    RECOMMENDATIONS: dict[str, list[str]] = {
        "Low":      [
            "Continue current study pace — student is on track.",
            "Encourage participation in advanced practice sessions.",
            "Maintain sleep schedule and avoid over-studying.",
        ],
        "Medium":   [
            "Schedule a bi-weekly progress check with assigned faculty.",
            "Review weak chapters identified in the SHAP analysis above.",
            "Suggest mock test retake strategy focused on low-scoring subjects.",
            "Share stress-relief resources and encourage peer study groups.",
        ],
        "High":     [
            "Arrange immediate one-on-one session with the faculty mentor.",
            "Review and adjust study plan to reduce subject overload.",
            "Monitor attendance daily for the next two weeks.",
            "Consider a temporary reduction in syllabus load.",
            "Connect student with peer mentor from top-performing cohort.",
        ],
        "Critical": [
            "URGENT: Schedule immediate counselling session.",
            "Notify parents and arrange a parent-faculty-student meeting.",
            "Refer student to psychological support / stress-management programme.",
            "Halt all additional test pressure for 2 weeks — focus on recovery.",
            "Daily check-ins by faculty until risk score stabilises below 60.",
        ],
    }

    recs = RECOMMENDATIONS.get(risk_level, RECOMMENDATIONS["Medium"])

    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(*TEXT)
    pdf.cell(0, 7, "Recommendations", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*TEXT)
    for rec in recs:
        pdf.set_x(22)
        pdf.multi_cell(170, 6, _safe(f">  {rec}"))
    pdf.ln(4)

    # ── 7. Footer ─────────────────────────────────────────────────────────────
    pdf.set_y(-18)
    pdf.set_font("Helvetica", "I", 7)
    pdf.set_text_color(*SLATE)
    pdf.set_draw_color(226, 232, 240)
    pdf.line(18, pdf.get_y(), 192, pdf.get_y())
    pdf.ln(2)
    pdf.cell(0, 4, _safe("CONFIDENTIAL - For institutional use only. Generated by JEE Dropout Prediction System."), align="C")

    return bytes(pdf.output())


# ═══════════════════════════════════════════════════════════════════════════════
# EXCEL — Batch Risk Summary
# ═══════════════════════════════════════════════════════════════════════════════

def generate_batch_excel(batch_id: int, students: list[Any]) -> bytes:
    """
    Generate an Excel workbook with:
      Sheet 1: Student Risk Summary — all students, score, level, top 3 factors
      Sheet 2: Batch Statistics    — avg score, distribution, top at-risk
    Returns raw .xlsx bytes. Raises RuntimeError if openpyxl not available.
    """
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()

    # ── Shared styles ─────────────────────────────────────────────────────────
    FILL_COLORS: dict[str, str] = {
        "Low":      "D1FAE5",
        "Medium":   "FEF3C7",
        "High":     "FEE2E2",
        "Critical": "EDE9FE",
    }
    TEXT_COLORS: dict[str, str] = {
        "Low":      "065F46",
        "Medium":   "92400E",
        "High":     "991B1B",
        "Critical": "4C1D95",
    }

    header_fill  = PatternFill("solid", fgColor="0F172A")
    header_font  = Font(bold=True, color="FFFFFF", size=10)
    title_font   = Font(bold=True, size=11, color="0F172A")
    thin_border  = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    def _header_row(ws, headers: list[str], row: int = 1):
        for col, label in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=label)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = thin_border

    def _auto_width(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    # ── Sheet 1: Student Risk Summary ─────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Student Risk Summary"
    ws1.row_dimensions[1].height = 22

    hdrs = ["#", "Name", "Student Code", "Email", "Risk Level", "Risk Score",
            "Factor 1", "Factor 2", "Factor 3", "Assessed At"]
    _header_row(ws1, hdrs)

    for row_idx, stu in enumerate(students, 2):
        latest = None
        if hasattr(stu, "risk_assessments") and stu.risk_assessments:
            latest = sorted(stu.risk_assessments, key=lambda r: r.assessed_at, reverse=True)[0]

        rl    = getattr(latest, "risk_level", "Unknown") if latest else "—"
        rs    = getattr(latest, "risk_score", None)      if latest else None
        at_dt = getattr(latest, "assessed_at", None)     if latest else None

        # Top 3 factors from ShapExplanation if available
        factors: list[str] = []
        if latest and hasattr(latest, "shap_explanations") and latest.shap_explanations:
            se   = latest.shap_explanations[-1]
            data = getattr(se, "top_factors", []) or []
            if isinstance(data, list):
                factors = [f.get("human_label", f.get("feature", "")) for f in data[:3]]

        row_data = [
            row_idx - 1,
            getattr(stu, "full_name",     "—"),
            getattr(stu, "student_code",  "—"),
            getattr(stu, "email",         "—"),
            rl,
            round(rs, 1) if rs is not None else "—",
            factors[0] if len(factors) > 0 else "—",
            factors[1] if len(factors) > 1 else "—",
            factors[2] if len(factors) > 2 else "—",
            at_dt.strftime("%d %b %Y %H:%M") if at_dt else "—",
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.border    = thin_border
            cell.alignment = Alignment(vertical="center")

            # Conditional fill on the Risk Level column (col 5)
            if col_idx == 5 and rl in FILL_COLORS:
                cell.fill = PatternFill("solid", fgColor=FILL_COLORS[rl])
                cell.font = Font(bold=True, color=TEXT_COLORS[rl])

    _auto_width(ws1)

    # ── Sheet 2: Batch Statistics ─────────────────────────────────────────────
    ws2 = wb.create_sheet(title="Batch Statistics")

    ws2["A1"] = f"Batch {batch_id} — Summary Statistics"
    ws2["A1"].font = title_font
    ws2.merge_cells("A1:D1")
    ws2.row_dimensions[1].height = 20

    assessments = []
    for stu in students:
        if hasattr(stu, "risk_assessments") and stu.risk_assessments:
            latest = sorted(stu.risk_assessments, key=lambda r: r.assessed_at, reverse=True)[0]
            assessments.append(latest)

    scores      = [a.risk_score for a in assessments if a.risk_score is not None]
    avg_score   = sum(scores) / len(scores) if scores else 0
    dist: dict[str, int] = {"Low": 0, "Medium": 0, "High": 0, "Critical": 0}
    for a in assessments:
        if a.risk_level in dist:
            dist[a.risk_level] += 1

    stats = [
        ("Total Students",          len(students)),
        ("Avg Risk Score",          f"{avg_score:.1f}"),
        ("Low Risk",                dist["Low"]),
        ("Medium Risk",             dist["Medium"]),
        ("High Risk",               dist["High"]),
        ("Critical Risk",           dist["Critical"]),
        ("High + Critical (Total)", dist["High"] + dist["Critical"]),
    ]

    for r_off, (label, value) in enumerate(stats, 3):
        ws2.cell(row=r_off, column=1, value=label).font = Font(bold=True, color="374151")
        ws2.cell(row=r_off, column=2, value=value)

    # Top at-risk sublist
    ws2.cell(row=11, column=1, value="Top At-Risk Students").font = title_font
    at_risk_hdrs = ["Name", "Code", "Risk Level", "Score"]
    _header_row(ws2, at_risk_hdrs, row=12)

    sorted_by_risk = sorted(
        [(stu, a) for stu, a in zip(students, assessments) if a.risk_level in ("High", "Critical")],
        key=lambda x: x[1].risk_score or 0,
        reverse=True,
    )[:10]

    for r_off, (stu, a) in enumerate(sorted_by_risk, 13):
        rl = a.risk_level
        row_vals = [getattr(stu, "full_name", "—"), getattr(stu, "student_code", "—"),
                    rl, round(a.risk_score or 0, 1)]
        for c_off, val in enumerate(row_vals, 1):
            cell = ws2.cell(row=r_off, column=c_off, value=val)
            cell.border = thin_border
            if c_off == 3 and rl in FILL_COLORS:
                cell.fill = PatternFill("solid", fgColor=FILL_COLORS[rl])
                cell.font = Font(bold=True, color=TEXT_COLORS[rl])

    _auto_width(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
